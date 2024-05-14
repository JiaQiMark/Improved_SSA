import numpy as np
from docx import Document
import matplotlib.pyplot as plt


def get_mean_std(prefix_path, prefix_path_myssa,
                 iterations,
                 func_index,
                 ssa, ssa_my, ssa_pso, ssa_sin, ssa_ic, ssa_wmr, gwo, pso, cso, ssamy_d100, ssamy_d500):

    epoch = 30
    max_iterations = 1000

    func_result_original = np.zeros((4, 4))
    func_result_improved_ssa = np.zeros((2, 4))
    func_result_only_weight = np.zeros((3, 2))

    func_result_sca_cssa = np.zeros((1, 2))

    load_data_ssa = np.zeros((epoch, max_iterations))
    load_data_myssa = np.zeros((epoch, max_iterations))
    load_data_ssa_sin = np.zeros((epoch, max_iterations))
    load_data_ssa_pso = np.zeros((epoch, max_iterations))
    load_data_ssa_ic = np.zeros((epoch, max_iterations))
    load_data_ssa_wmr = np.zeros((epoch, max_iterations))
    load_data_gwo = np.zeros((epoch, max_iterations))
    load_data_pso = np.zeros((epoch, max_iterations))
    load_data_cso = np.zeros((epoch, max_iterations))

    if ssa_ic:
        load_data_ssa_ic = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_ssa_ic.npy")
        last_row = load_data_ssa_ic[:, iterations]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)

        func_result_only_weight[0][0] = mean_v
        func_result_only_weight[0][1] = std_v
    if ssa_wmr:
        load_data_ssa_wmr = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_ssa_wmr.npy")
        last_row = load_data_ssa_wmr[:, iterations]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)

        func_result_only_weight[1][0] = mean_v
        func_result_only_weight[1][1] = std_v
    if ssa:
        load_data_ssa = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_ssa.npy")
        selected_row = load_data_ssa[:, iterations[0]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[0][0] = mean_v
        func_result_original[0][1] = std_v
        # print(f"-----------f{func_index} SSA------------     Mean = {mean_v}   Std = {std_v}")
        selected_row = load_data_ssa[:, iterations[1]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[0][2] = mean_v
        func_result_original[0][3] = std_v

    if gwo:
        load_data_gwo = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_gwo.npy")
        selected_row = load_data_gwo[:, iterations[0]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[1][0] = mean_v
        func_result_original[1][1] = std_v
        # print(f"-----------f{func_index} GWO------------    Mean = {mean_v}   Std = {std_v}")
        selected_row = load_data_gwo[:, iterations[1]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[1][2] = mean_v
        func_result_original[1][3] = std_v

    if pso:
        load_data_pso = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_pso.npy")
        selected_row = load_data_pso[:, iterations[0]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[2][0] = mean_v
        func_result_original[2][1] = std_v
        # print(f"-----------f{func_index} PSO------------    Mean = {mean_v}   Std = {std_v}")
        selected_row = load_data_pso[:, iterations[1]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[2][2] = mean_v
        func_result_original[2][3] = std_v

    if cso:
        load_data_cso = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_cso.npy")

    if ssa_pso:
        load_data_ssa_pso = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_ssa_pso.npy")
        last_row = load_data_ssa_pso[:, iterations[0]]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)
        func_result_improved_ssa[0][0] = mean_v
        func_result_improved_ssa[0][1] = std_v

        last_row = load_data_ssa_pso[:, iterations[1]]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)
        func_result_improved_ssa[0][2] = mean_v
        func_result_improved_ssa[0][3] = std_v
    if ssa_sin:
        load_data_ssa_sin = np.load(f"{prefix_path}\\f{func_index}_{epoch}_{max_iterations}_ssa_sin.npy")
        last_row = load_data_ssa_sin[:, iterations[0]]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)
        func_result_improved_ssa[1][0] = mean_v
        func_result_improved_ssa[1][1] = std_v

        last_row = load_data_ssa_sin[:, iterations[1]]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)
        func_result_improved_ssa[1][2] = mean_v
        func_result_improved_ssa[1][3] = std_v

    if ssa_my:
        load_data_myssa = np.load(f"{prefix_path_myssa}\\f{func_index}_{epoch}_{max_iterations}_myssa.npy")

        selected_row = load_data_myssa[:, iterations[0]]
        mean_v = np.mean(selected_row)
        # np.sum(np.log10(selected_row + 1e-320))/30
        std_v = np.std(selected_row)
        func_result_original[3][0] = mean_v
        func_result_original[3][1] = std_v
        func_result_only_weight[2][0] = mean_v
        func_result_only_weight[2][1] = std_v
        # print(f"-----------f{func_index} CIW-SSA ------------    Mean = {mean_v}   Std = {std_v}")
        selected_row = load_data_myssa[:, iterations[1]]
        mean_v = np.mean(selected_row)
        std_v = np.std(selected_row)
        func_result_original[3][2] = mean_v
        func_result_original[3][3] = std_v

    solution_result_100_500 = np.zeros((2, 2))
    if ssamy_d100:
        load_data_myssa = np.load(f"result_7_hd100_myssa\\f{func_index}_{epoch}_{max_iterations}_myssa.npy")
        last_row = load_data_myssa[:, iterations]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)
        solution_result_100_500[0][0] = mean_v
        solution_result_100_500[0][1] = std_v
    if ssamy_d500:
        load_data_myssa = np.load(f"result_8_hd500_myssa\\f{func_index}_{epoch}_{max_iterations}_myssa.npy")
        last_row = load_data_myssa[:, iterations]
        mean_v = np.mean(last_row)
        std_v = np.std(last_row)
        solution_result_100_500[1][0] = mean_v
        solution_result_100_500[1][1] = std_v

    return func_result_improved_ssa
    # return func_result_original


pref_path = "result_3_low_dimension_30"
pref_path_myssa = "result_6_ld30_myssa"
pref_path_save = "docx_mean_std\\02_improved_200_1000_iterations.docx"


function_index_list = [0,
                       1,   2,  3,  4,  5,  6,  7,  8,  9, 10,
                       11, 12, 13, 14, 15, 16, 17, 18, 19, 20]

# function_index_list = [0,
#                        1,   2,  3,  4,  6,  7,  8,  9, 10,
#                        11, 12, 13, 14, 15]

doc = Document()

# for i in range(9, 10):
for i in range(1, len(function_index_list)):
    if i in [14, 18]:
        continue
    function_index = function_index_list[i]

    f_ssa_ic = 0
    f_ssa_wmr = 0

    f_ssa = 0
    f_gwo = 0
    f_pso = 0
    f_cso = 0

    f_ssa_pso = 1
    f_ssa_sin = 1

    f_ssa_my = 0

    f_ssamy_d100 = 0
    f_ssamy_d500 = 0

    bms_result = get_mean_std(pref_path, pref_path_myssa,
                              [200, -1],
                              function_index,
                              f_ssa, f_ssa_my, f_ssa_pso, f_ssa_sin, f_ssa_ic, f_ssa_wmr, f_gwo, f_pso, f_cso,
                              f_ssamy_d100, f_ssamy_d500)
    bms_result_t = bms_result.T
    bms_result_t = np.array(bms_result_t)
    rows, cols = bms_result_t.shape
    # 插入第一个表格
    table1 = doc.add_table(rows, cols)

    for row in range(rows):
        for col in range(cols):
            table1.cell(row, col).text = "%.4E" % bms_result_t[row, col]

doc.save(pref_path_save)



